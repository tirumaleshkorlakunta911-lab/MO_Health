from openpyxl import load_workbook

import config


def load_records():

    workbook = load_workbook(
        config.EXCEL_PATH
    )

    worksheet = workbook.active

    headers = {}

    for cell in worksheet[1]:

        if cell.value is not None:

            headers[
                str(cell.value).strip()
            ] = cell.column

    required = [
        config.COL_NAME,
        config.COL_RBS,
        config.COL_BP_SYSTOLIC,
        config.COL_BP_DIASTOLIC,
        config.COL_PULSE,
        config.COL_STATUS,
    ]

    for column in required:

        if column not in headers:

            raise ValueError(
                f"Missing Excel column: {column}"
            )

    records = []

    for row_number in range(
        2,
        worksheet.max_row + 1,
    ):

        name = worksheet.cell(
            row_number,
            headers[config.COL_NAME],
        ).value

        if not name:
            continue

        status = worksheet.cell(
            row_number,
            headers[config.COL_STATUS],
        ).value

        status = str(status or "").strip()

        if status == config.STATUS_SUBMITTED:
            continue

        records.append(
            {
                "row_number": row_number,

                config.COL_NAME:
                    worksheet.cell(
                        row_number,
                        headers[config.COL_NAME],
                    ).value,

                config.COL_RBS:
                    worksheet.cell(
                        row_number,
                        headers[config.COL_RBS],
                    ).value,

                config.COL_BP_SYSTOLIC:
                    worksheet.cell(
                        row_number,
                        headers[config.COL_BP_SYSTOLIC],
                    ).value,

                config.COL_BP_DIASTOLIC:
                    worksheet.cell(
                        row_number,
                        headers[config.COL_BP_DIASTOLIC],
                    ).value,

                config.COL_PULSE:
                    worksheet.cell(
                        row_number,
                        headers[config.COL_PULSE],
                    ).value,
            }
        )

    return records


def update_status(
    row_number,
    status,
):

    workbook = load_workbook(
        config.EXCEL_PATH
    )

    worksheet = workbook.active

    status_column = None

    for cell in worksheet[1]:

        if (
            str(cell.value or "").strip()
            == config.COL_STATUS
        ):

            status_column = cell.column
            break

    if status_column is None:

        raise ValueError(
            "Status column not found."
        )

    worksheet.cell(
        row=row_number,
        column=status_column,
        value=status,
    )

    workbook.save(
        config.EXCEL_PATH
    )


def mark_submitted(row_number):

    update_status(
        row_number,
        config.STATUS_SUBMITTED,
    )


def mark_failed(
    row_number,
    reason,
):

    update_status(
        row_number,
        f"{config.STATUS_FAILED}: {reason}",
    )